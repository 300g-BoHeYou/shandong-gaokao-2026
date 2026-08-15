# -*- coding: utf-8 -*-
"""
生成《试验班查证清单.xlsx》: 312个缺少具体专业的试验班志愿类, 按院校分组。

列: 序号 | 院校代码 | 院校名称 | 志愿类全名 | 备注 | 选科要求 | 最低分 | 最低位次 | 招生类型
    | 查证状态 | 包含专业(待填,顿号分隔) | 来源链接/年份(待填) | 备注(待填)
"""
import re
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "2026山东高考第一批次投档.xlsx"
OUT = ROOT / "试验班查证清单.xlsx"

df = pd.read_excel(XLSX, sheet_name="第一次投档表", dtype=str)


def parse_name(name):
    m = re.match(r"^([0-9A-Za-z]{1,4})\s+(.*)$", str(name).strip())
    return (m.group(1), m.group(2).strip()) if m else ("", str(name).strip())


def split_core_note(full):
    m = re.search(r"[（(]", full)
    if m:
        return full[: m.start()].strip(), full[m.start() + 1:].rstrip(")）").strip()
    return full, ""


def school_parts(s):
    m = re.match(r"^([0-9A-Za-z]+)(.*)$", str(s).strip())
    if m and m.group(1) and m.group(2):
        return m.group(1), m.group(2).strip()
    return "", str(s).strip()


TEST_BAN = re.compile(r"试验班|实验班|试点班")
rows = []
for _, r in df.iterrows():
    code, full = parse_name(r["专业代号及名称"])
    core, note = split_core_note(full)
    if not TEST_BAN.search(core):
        continue
    sc, school = school_parts(r["院校代号及名称"])
    rows.append({
        "sc": sc, "school": school, "full": full, "core": core, "note": note,
        "xk": str(r["选考科目要求"]).strip(),
        "score": int(float(r["最低分"])) if str(r["最低分"]).replace(".", "", 1).isdigit() else None,
        "rank": int(float(r["最低位次"])) if str(r["最低位次"]).replace(".", "", 1).isdigit() else None,
        "typ": str(r["专项计划/中外合作"]).strip() if pd.notna(r["专项计划/中外合作"]) else "普通",
    })

# 按院校名称分组, 校内按位次排序
rows.sort(key=lambda r: (r["school"], r["rank"] if r["rank"] is not None else 10**9, r["full"]))
print("试验班志愿类总数:", len(rows))
print("涉及院校数:", len({r["school"] for r in rows}))

# 同校同全名重复检查
units = Counter((r["school"], r["full"]) for r in rows)
dups = {k: v for k, v in units.items() if v > 1}
if dups:
    print("同校同全名重复(需人工确认是否普通+专项并存):")
    for k, v in dups.items():
        print("  ", v, "×", k)

# ---------- 写 Excel ----------
wb = Workbook()
ws = wb.active
ws.title = "试验班查证清单"
headers = ["序号", "院校代码", "院校名称", "志愿类全名", "备注", "选科要求",
           "最低分", "最低位次", "招生类型", "查证状态", "包含专业(顿号分隔,待填)",
           "来源链接/年份(待填)", "备注(待填,如图片形式)"]
ws.append(headers)
head_fill = PatternFill("solid", fgColor="1D4ED8")
head_font = Font(color="FFFFFF", bold=True)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

school_fill = PatternFill("solid", fgColor="EFF6FF")
prev_school = None
for i, r in enumerate(rows, 1):
    ws.append([i, r["sc"], r["school"], r["full"], r["note"], r["xk"],
               r["score"], r["rank"], r["typ"], "待填", "", "", ""])
    rr = ws.max_row
    if r["school"] != prev_school:
        for c in range(1, len(headers) + 1):
            ws.cell(row=rr, column=c).fill = school_fill  # 院校分组底色
        prev_school = r["school"]

widths = [6, 9, 20, 34, 22, 11, 8, 9, 12, 9, 40, 26, 20]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:M{ws.max_row}"
wb.save(OUT)
print("已生成:", OUT)

# ---------- 打印对话框清单 ----------
print("\n===== 需查证清单(按院校分组) =====")
cur = None
for i, r in enumerate(rows, 1):
    if r["school"] != cur:
        cur = r["school"]
        cnt = sum(1 for x in rows if x["school"] == cur)
        print(f"\n【{cur}】{cnt} 个")
    note = f"（{r['note']}）" if r["note"] else ""
    xk = r["xk"]
    sc_ = f"{r['score']}分" if r["score"] else "?"
    rk_ = f"{r['rank']}名" if r["rank"] else "?"
    typ_ = "" if r["typ"] == "普通" else f"[{r['typ']}]"
    print(f"  {r['sc']} {r['full']}{note} ｜ {xk} ｜ {sc_}/{rk_} {typ_}")
