# -*- coding: utf-8 -*-
"""
生成《大类查证清单（软科未收录）.xlsx》: 软科未收录院校的大类招生行。
本次范围: 国防科技大学、中国社会科学院大学 (+后续可扩展到其余未收录院校)。

说明: 中国科学技术大学已被软科收录(自动筛减), 中国科学院大学无大类行,
      陆军工程大学/军医大学等均为具体专业招生, 均不在此表。
"""
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "2026山东高考第一批次投档.xlsx"
OUT = ROOT / "大类查证清单（软科未收录）.xlsx"

with open(ROOT / "软科中国大学专业排名_2026.json", encoding="utf-8") as f:
    ruanke = json.load(f)
rk_schools = {s for m in ruanke.values() for cls in m.values() for rks in cls.values() for s in rks}

with open(ROOT / "scripts" / "catalog_tree.json", encoding="utf-8") as f:
    cat = json.load(f)
lei_names = set()
for cls in cat["tree"].values():
    for lei in cls:
        lei_names.add(lei)

td = pd.read_excel(XLSX, sheet_name="第一次投档表", dtype=str)


def parse_name(name):
    m = re.match(r"^([0-9A-Za-z]{1,4})\s+(.*)$", str(name).strip())
    return (m.group(1), m.group(2).strip()) if m else ("", str(name).strip())


def split_core_note(full):
    m = re.search(r"[（(]", full)
    if m:
        return full[: m.start()].strip(), full[m.start() + 1:].rstrip(")）").strip()
    return full, ""


def school_parts(s):
    m = re.match(r"^([0-9A-Za-z]+)(.*)$", str(s).strip())
    if m and m.group(1) and m.group(2):
        return m.group(1), m.group(2).strip()
    return "", str(s).strip()


# 本次范围
FOCUS = {"国防科技大学", "中国社会科学院大学"}
rows = []
for _, r in td.iterrows():
    _c, full = parse_name(r["专业代号及名称"])
    core, note = split_core_note(full)
    sc, school = school_parts(r["院校代号及名称"])
    if school not in FOCUS:
        continue
    if school in rk_schools:
        continue  # 已收录
    if core not in lei_names:
        continue  # 只要大类行
    rows.append({
        "sc": sc, "school": school, "full": full, "note": note,
        "xk": str(r["选考科目要求"]).strip(),
        "score": int(float(r["最低分"])),
        "rank": int(float(r["最低位次"])),
        "typ": str(r["专项计划/中外合作"]).strip() if pd.notna(r["专项计划/中外合作"]) else "普通",
    })

rows.sort(key=lambda r: (r["school"], r["rank"]))
print("本次查证行数:", len(rows))

# ---------- Excel ----------
wb = Workbook()
ws = wb.active
ws.title = "大类查证清单"
headers = ["序号", "院校代码", "院校名称", "志愿类全名", "备注", "选科要求",
           "最低分", "最低位次", "招生类型", "查证状态", "包含专业(顿号分隔,待填)",
           "来源链接/年份(待填)", "备注(待填,如图片形式)"]
ws.append(headers)
head_fill = PatternFill("solid", fgColor="1D4ED8")
head_font = Font(color="FFFFFF", bold=True)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for i, r in enumerate(rows, 1):
    ws.append([i, r["sc"], r["school"], r["full"], r["note"], r["xk"],
               r["score"], r["rank"], r["typ"], "待填", "", "", ""])

widths = [6, 9, 20, 34, 22, 11, 8, 9, 12, 9, 40, 26, 20]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:M{ws.max_row}"
wb.save(OUT)
print("已生成:", OUT)

# ---------- 对话框清单 ----------
print("\n===== 需查证清单 =====")
cur = None
for i, r in enumerate(rows, 1):
    if r["school"] != cur:
        cur = r["school"]
        print(f"\n【{cur}】")
    note = f"（{r['note']}）" if r["note"] else ""
    print(f"  {r['sc']} {r['full']}{note} ｜ {r['xk']} ｜ {r['score']}分/{r['rank']}名")
