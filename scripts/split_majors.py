# -*- coding: utf-8 -*-
"""把 Sheet2 的「专业类」按其中每个专业拆成独立行, 并逐专业计算:
   - 软科2026排名(该校该专业)
   - 对应一批次(普通批)参考分/位次(优先同名专业, 否则含该专业的大类/已查证试验班)
   - 名称校验: 目录中是否存在(疑似错字自动修正) + 该校是否有开设迹象(无则标注)

输出: 强基查证清单.xlsx(原表, Sheet1 官方字段 + Sheet2 拆分后) 与 校验报告。
"""
import difflib
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
MERGED = ROOT / "强基查证清单_官方核对.xlsx"
SRC = MERGED if MERGED.exists() else ROOT / "强基查证清单.xlsx"
if len(sys.argv) > 2:
    SRC = Path(sys.argv[2])
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "强基查证清单.xlsx"
EXCL = re.compile(r"分校|校区|马来西亚|中外合作|合作办学|定向|预科|走读|联合培养|专项")

# ---------- 目录 ----------
with open(ROOT / "scripts" / "catalog_tree.json", encoding="utf-8") as f:
    cat = json.load(f)
LEI2MAJOR, ALL_MAJOR, MAJOR2MEN = {}, set(), {}
for _men, cls in cat["tree"].items():
    for lei, ms in cls.items():
        LEI2MAJOR[lei] = [n for _c, n in ms]
        for _c, n in ms:
            ALL_MAJOR.add(n)
            MAJOR2MEN[n] = _men
ALL_LEI = set(LEI2MAJOR)

# ---------- 软科 ----------
with open(ROOT / "软科中国大学专业排名_2026.json", encoding="utf-8") as f:
    rk = json.load(f)
RK_MAJOR, NAME2KEYS, SPECIAL = {}, {}, set()
for men, cls in rk.items():
    for lei, ms in cls.items():
        for code_name, ranks in ms.items():
            nm = code_name.split(" ", 1)[1] if " " in code_name else code_name
            RK_MAJOR[code_name] = dict(ranks)
            NAME2KEYS.setdefault(nm, []).append(code_name)
            if ranks and set(ranks.values()) == {1}:
                SPECIAL.add(code_name)
RK_SCHOOLS = {s for m in RK_MAJOR.values() for s in m}
RK_NORM = {}
for s in RK_SCHOOLS:
    RK_NORM.setdefault(re.sub(r"\s+", "", s).replace("（", "(").replace("）", ")"), s)
SUF = [r"克拉玛依校区$", r"秦皇岛分校$", r"威海分校$", r"青岛校区$", r"济南校区$", r"医学部$", r"医学院$",
       r"[（(][^（()）]*校区[）)]$", r"[（(][^（()）]*分校[）)]$", r"[（(][^（()）]*[）)]$"]


def canon(name):
    n = re.sub(r"\s+", "", name).replace("（", "(").replace("）", ")")
    if n in RK_NORM:
        return RK_NORM[n]
    cur = n
    for _ in range(4):
        nxt = cur
        for p in SUF:
            c = re.sub(p, "", nxt)
            if c != nxt and c in RK_NORM:
                return RK_NORM[c]
        if nxt == cur:
            break
        cur = nxt
    return ""


def rank_info(school, major):
    """{'rk','st','cnt'} st: ranked/unranked/special/notpub/nosc"""
    rs = canon(school)
    keys = NAME2KEYS.get(major) or []
    if not keys:
        return {"rk": None, "st": "notpub", "cnt": 0}
    key = keys[0]
    if key in SPECIAL:
        return {"rk": None, "st": "special", "cnt": len(RK_MAJOR[key]),
                "inList": bool(rs and rs in RK_MAJOR[key])}
    ranks = RK_MAJOR[key]
    if not rs:
        return {"rk": None, "st": "nosc", "cnt": len(ranks)}
    if rs in ranks:
        return {"rk": ranks[rs], "st": "ranked", "cnt": len(ranks)}
    return {"rk": None, "st": "unranked", "cnt": len(ranks)}


# ---------- 投档表 ----------
td = pd.read_excel(ROOT / "2026山东高考第一批次投档.xlsx", sheet_name="第一次投档表", dtype=str)


def split_core(full):
    m = re.search(r"[（(]", str(full))
    return (str(full)[: m.start()].strip(), str(full)[m.start() + 1:].rstrip(")）").strip()) if m else (str(full).strip(), "")


def school_of(s):
    m = re.match(r"^([0-9A-Za-z]+)(.*)$", str(s).strip())
    return m.group(2).strip() if m and m.group(1) and m.group(2) else str(s).strip()


mc = json.loads((ROOT / "scripts" / "manual_check.json").read_text(encoding="utf-8")) if (ROOT / "scripts" / "manual_check.json").exists() else {}
SYB = mc.get("试验班", {})

# 各校强基简章全文(用于判断"该校是否提及该专业")
CORPUS = {}
_cp = ROOT / "qj_cache" / "corpus.json"
if _cp.exists():
    _c = json.loads(_cp.read_text(encoding="utf-8"))
    CORPUS = {k: re.sub(r"\s+", "", v.get("text", "")) for k, v in _c.items()}
for _n, _f in [("中央民族大学", "fix_69437_600.html"), ("上海交通大学", "fix_38987_52.html"),
               ("四川大学", "fix_scu_jz.html"), ("兰州大学", "lzu_jz.html")]:
    _p = ROOT / "qj_cache" / _f
    if _p.exists() and len(CORPUS.get(_n, "")) < 1200:
        CORPUS[_n] = re.sub(r"\s+", "", re.sub(r"<[^>]+>", " ", _p.read_text(encoding="utf-8", errors="replace")))

# 固定名称对照(经核对); 以及经确认属于"抄错"应删除的条目
FIX_MAP = {
    "土木水利与海洋工程": "土木、水利与海洋工程",
    "通讯工程": "通信工程",
    "微电子学": "微电子科学与工程",
}
MANUAL_FIX = {("清华大学", "微电子学"): "微电子科学与工程"}
MANUAL_DROP = {
    ("东南大学", "材料物理"),
    ("东南大学", "系统科学"),
    ("清华大学", "伦理学"),
    ("清华大学", "逻辑学"),
    ("清华大学", "宗教学"),
    ("北京大学", "语言学"),
}
SUGGEST = {}

school_rows = {}
for _, r in td.iterrows():
    full = re.sub(r"^[0-9A-Za-z]{1,4}\s+", "", str(r["专业代号及名称"]).strip())
    core, _n = split_core(full)
    sch = school_of(r["院校代号及名称"])
    typ = str(r["专项计划/中外合作"]).strip() if pd.notna(r["专项计划/中外合作"]) else "普通"
    try:
        sc, rk_ = int(float(r["最低分"])), int(float(r["最低位次"]))
    except (TypeError, ValueError):
        continue
    code = re.match(r"^([0-9A-Za-z]+)", str(r["院校代号及名称"]).strip())
    key = f"{code.group(1) if code else ''}|{full}|{rk_}"
    e = SYB.get(key) or {}
    syb = e.get("majors", []) if e.get("mode") == "list" else []
    school_rows.setdefault(sch, []).append((core, sc, rk_, typ, syb, full))


TSINGHUA_LI = ["数学与应用数学", "物理学", "化学", "生物科学", "信息与计算科学",
               "数理基础科学", "化学生物学", "理论与应用力学"]
TSINGHUA_WEN = ["中国语言文学类（古文字学方向）", "中国语言文学类（古文字学）", "历史学类", "哲学类"]


# 门类 -> 科类试验班关键词(该校按试验班招生时, 用对应科类试验班的最低分作参考)
MEN2SYB = {
    "理学": ["理科试验班", "自然科学试验班"],
    "工学": ["工科试验班", "技术科学试验班"],
    "哲学": ["文科试验班", "人文科学试验班"],
    "文学": ["文科试验班", "人文科学试验班"],
    "历史学": ["文科试验班", "人文科学试验班"],
    "经济学": ["经济管理试验班", "社会科学试验班"],
    "管理学": ["经济管理试验班", "社会科学试验班"],
    "法学": ["法学试验班", "社会科学试验班"],
    "医学": ["医学试验班", "自然科学试验班"],
    "农学": ["自然科学试验班"],
    "教育学": ["社会科学试验班"],
    "艺术学": ["艺术试验班"],
}


def _pool(school):
    """候选志愿类池: 本校 + 同一大学的医学部/医学院(如"北京大学医学部")"""
    out = list(school_rows.get(school, []))
    for sib, rows in school_rows.items():
        if sib != school and sib.startswith(school) and re.search(r"医学部|医学院|医学中心", sib):
            out.extend(rows)
    return [(c, sc, rk_, typ, syb, full) for c, sc, rk_, typ, syb, full in out
            if typ == "普通" and not EXCL.search(full)]


def ref_of(school, major):
    """分级取参考(同级内取最低分, 避免创新班等偏高):
       1) 同名专业志愿类  2) 目录含该专业的大类  3) 已人工查证含该专业的试验班
       4) 按门类对应的科类试验班(如理科试验班/工科试验班/人文科学试验班)
    返回 (分, 位次, 参考志愿类, 口径)。"""
    rows = _pool(school)

    def pick(pred):
        cand = [(sc, rk_, c) for c, sc, rk_, _t, _y, _f in rows if pred(c, _y)]
        return min(cand) if cand else None

    hit = pick(lambda c, _y: c == major)
    if hit:
        return (hit[0], hit[1], hit[2], "同名专业")
    hit = pick(lambda c, _y: c in LEI2MAJOR and major in LEI2MAJOR[c])
    if hit:
        return (hit[0], hit[1], hit[2], "大类")
    hit = pick(lambda c, y: major in y)
    if hit:
        return (hit[0], hit[1], hit[2], "试验班")
    men = MAJOR2MEN.get(major, "")
    kws = list(MEN2SYB.get(men, []))
    if men in ("理学", "工学"):
        kws += ["理科试验班", "工科试验班", "自然科学试验班"]   # 理工互通回退
    elif men in ("医学", "农学"):
        kws += ["自然科学试验班", "医科试验班"]
    for kw in kws:
        hit = pick(lambda c, _y, kw=kw: c.startswith(kw) or kw in c)
        if hit:
            return (hit[0], hit[1], hit[2], "试验班(按科类)")
    return (None, None, "", "")


def norm_major(n, school=""):
    """名称规范化 + 错字纠正。返回 (名称, 警告)
    仅按固定对照表/高相似度修正; 其余只给提示, 不擅自改名。"""
    raw = n.strip()
    if school and (school, raw) in MANUAL_FIX:
        return MANUAL_FIX[(school, raw)], ""
    x = re.sub(r"^含", "", raw)
    x = re.sub(r"专业$", "", x)
    x = re.sub(r"\s+", "", x)
    if x in ALL_MAJOR:
        return x, ""
    if x in ALL_LEI:
        return x, "这是大类名，不是具体专业"
    if x in FIX_MAP:
        return FIX_MAP[x], ""
    if x in SUGGEST:
        return x, f"2026目录中无“{x}”，疑似应为「{SUGGEST[x]}」，请核对"
    close = difflib.get_close_matches(x, list(ALL_MAJOR), n=2, cutoff=0.9)
    if close and abs(len(close[0]) - len(x)) <= 4:
        return close[0], ""
    sug = difflib.get_close_matches(x, list(ALL_MAJOR), n=1, cutoff=0.6)
    hint = f"，疑似应为「{sug[0]}」" if sug else ""
    return x, f"2026本科专业目录中无此专业名{hint}，请核对"


def main():
    wb = load_workbook(SRC)
    ws1, ws2 = wb["院校信息"], wb["专业信息"]
    hdr = [c.value for c in ws2[1]]
    ixc = {h: i + 1 for i, h in enumerate(hdr)}
    already = "包含专业" in hdr          # 已是拆分后的表(再次运行)
    inc_col = "包含专业" if already else "包含专业（顿号分隔）"
    note_col = "备注" if "备注" in hdr else "备注（括号内方向）"
    rows = []
    last_lei, last_subj, last_note = "", "", ""
    for r in range(2, ws2.max_row + 1):
        vals = {h: ws2.cell(row=r, column=i).value for h, i in ixc.items()}
        school = str(vals.get("院校名称") or "").strip()
        if not school:
            continue
        # 合并单元格: 非首行读到的类名/科目/备注为空, 向下填充
        lei = str(vals.get("强基专业类名称") or "").strip()
        subj = str(vals.get("考查科目") or "").strip()
        note = str(vals.get(note_col) or "").strip()
        if lei:
            last_lei = lei
        else:
            lei = last_lei
        if subj:
            last_subj = subj
        else:
            subj = last_subj
        if note:
            last_note = note
        else:
            note = last_note
        inc = str(vals.get(inc_col) or "").strip()
        parts = [inc] if already else ([p for p in inc.split("、") if p.strip()] or [""])
        for p in parts:
            rows.append(dict(school=school, lei=lei, raw=str(p).strip(), subj=subj, note=note))

    report, out = [], []
    for i, r in enumerate(rows, 1):
        nm, warn = norm_major(r["raw"], r["school"]) if r["raw"] else ("", "")
        if (r["school"], r["raw"]) in MANUAL_DROP or (r["school"], nm) in MANUAL_DROP:
            report.append((r["school"], r["lei"], r["raw"], "已删除(抄写有误)"))
            continue
        ri = rank_info(r["school"], nm) if nm else {}
        sc, rk_, refcat, kind = ref_of(r["school"], nm) if nm else (None, None, "", "")
        # 开设依据(严格): 软科上榜 / 特色专业且在开设名单内 / 一批次直接招生 / 该校强基简章提及
        # 注: 软科"未上榜"不算依据(用户明确: 不接受"专业水平低故未收录"这一解释)
        evidence = False
        if ri.get("st") == "ranked" or (ri.get("st") == "special" and ri.get("inList")):
            evidence = True
        if any(core == nm for core, _s, _r, typ, _y, full in school_rows.get(r["school"], []) if typ == "普通"):
            evidence = True
        if nm and nm in CORPUS.get(r["school"], ""):
            evidence = True
        warn2 = warn
        if nm and not warn and not evidence:
            warn2 = "未能证实该校开设该专业（简章未提及、一批次无直接招生、软科未上榜）"
        out.append(dict(院校名称=r["school"], 强基专业类名称=r["lei"], 包含专业=nm,
                        软科排名=(ri.get("rk") if ri.get("st") == "ranked" else ""),
                        软科状态={"ranked": "上榜", "unranked": "未上榜", "special": "特色专业",
                                  "notpub": "未发布", "nosc": "该校未收录软科"}.get(ri.get("st", ""), ""),
                        对应普通批最低分=sc, 对应普通批最低位次=rk_, 参考志愿类=refcat, 参考口径=kind,
                        考查科目=r["subj"], 备注=r["note"], 校验=warn2))
        if warn2:
            report.append((r["school"], r["lei"], nm, warn2))

    # 参考分补齐: 同一专业类内若无该专业的直接参考, 沿用该类其他专业的最低分参考(标注"同专业类")
    df = pd.DataFrame(out)
    filled = 0
    for (sch, lei), g in df.groupby(["院校名称", "强基专业类名称"], sort=False):
        have = g[g["对应普通批最低位次"].notna()]
        if have.empty:
            continue
        s_best = have["对应普通批最低分"].min()
        row_best = have[have["对应普通批最低分"] == s_best].iloc[0]
        miss = g[g["对应普通批最低位次"].isna()].index
        for idx in miss:
            df.at[idx, "对应普通批最低分"] = row_best["对应普通批最低分"]
            df.at[idx, "对应普通批最低位次"] = row_best["对应普通批最低位次"]
            df.at[idx, "参考志愿类"] = row_best["参考志愿类"]
            df.at[idx, "参考口径"] = "同专业类"
            filled += 1
    df.insert(0, "序号", range(1, len(df) + 1))

    print("拆分后专业行:", len(df), "(输入专业行:", len(rows), ")")
    print("\n== 处理记录 ==")
    for x in report:
        print("  " + " | ".join(str(v) for v in x))
    print(f"\n参考分补齐(同专业类): {filled} 条")
    print("有软科排名的专业行:", (df["软科排名"] != "").sum(),
          " 有普通批位次的:", df["对应普通批最低位次"].notna().sum())

    if OUT != SRC or True:
        # 写回(Sheet1 保留, Sheet2 替换为拆分表)
        if "专业信息" in wb.sheetnames:
            del wb["专业信息"]
        ws = wb.create_sheet("专业信息")
        cols = list(df.columns)
        ws.append(cols)
        for rec in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in rec])
        fill = PatternFill("solid", fgColor="1D4ED8")
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        widths = [5, 14, 26, 22, 10, 10, 14, 16, 22, 12, 20, 18, 26]
        for c, w in enumerate(widths[:len(cols)], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = "A2"

        # 同一(院校, 专业类)的类级列纵向合并单元格
        merge_cols = [cols.index(x) + 1 for x in ("强基专业类名称", "考查科目", "备注") if x in cols]
        r = 2
        while r <= ws.max_row:
            sch = ws.cell(row=r, column=cols.index("院校名称") + 1).value
            lei = ws.cell(row=r, column=cols.index("强基专业类名称") + 1).value
            r2 = r
            while (r2 + 1 <= ws.max_row
                   and ws.cell(row=r2 + 1, column=cols.index("院校名称") + 1).value == sch
                   and ws.cell(row=r2 + 1, column=cols.index("强基专业类名称") + 1).value == lei):
                r2 += 1
            if r2 > r:
                for c in merge_cols:
                    ws.merge_cells(start_row=r, start_column=c, end_row=r2, end_column=c)
            for c in merge_cols:
                ws.cell(row=r, column=c).alignment = Alignment(vertical="center", wrap_text=True)
            r = r2 + 1
        wb.save(OUT)
        print("\n已保存:", OUT, "(类专业列已纵向合并)")


if __name__ == "__main__":
    main()
