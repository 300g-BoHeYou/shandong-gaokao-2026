# -*- coding: utf-8 -*-
"""用 openpyxl 直接读取齐鲁师范学院'数据计算及应用'行的原始单元格值。"""
import openpyxl

wb = openpyxl.load_workbook("2026山东高考第一批次投档.xlsx", read_only=False, data_only=True)
ws = wb["第一次投档表"]
print("sheet 行数(含表头):", ws.max_row, "列数:", ws.max_column)

# 打印表头
hdr = [c.value for c in ws[1]]
print("表头:", hdr)

found = []
for row in ws.iter_rows(min_row=2):
    v = row[0].value
    if v is not None and "数据计算及应用" in str(v):
        found.append([c.value for c in row])

print("匹配行数:", len(found))
for r in found:
    for i, cell in enumerate(r):
        print(f"  col{i}[{hdr[i] if i < len(hdr) else '?'}]: len={len(str(cell)) if cell is not None else 0} | {repr(cell)}")

# 查看 sharedStrings 中是否有截断
print("\n== 检查该单元格在 sharedStrings 中的原始内容 ==")
for r in found:
    v = r[0]
    if isinstance(v, str) and "数据计算及应用" in v:
        print("repr:", repr(v))
        print("尾20字符:", repr(v[-20:]))
        print("是否以')'结尾:", v.endswith(")"))
