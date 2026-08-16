# -*- coding: utf-8 -*-
"""生成《大类查证清单（剩余未收录）.xlsx》: 校区回退后真正未收录且有大类行的院校(重庆外语外事学院)。"""
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent.parent
XLSX = ROOT / "2026山东高考第一批次投档.xlsx"
OUT = ROOT / "大类查证清单（剩余未收录）.xlsx"

td = pd.read_excel(XLSX, sheet_name="第一次投档表", dtype=str)


def parse_name(name):
    m = re.match(r"^([0-9A-Za-z]{1,4})\s+(.*)$", str(name).strip())
    return (m.group(1), m.group(2).strip()) if m else ("", str(name).strip())


def split_core_note(full):
    m = re.search(r"[（(]", full)
    if m:
        return full[: m.start()].strip(), full[m.start() + 1:].rstrip(")）").strip()
    return full, ""


def school_parts(s):
    m = re.match(r"^([0-9A-Za-z]+)(.*)$", str(s).strip())
    if m and m.group(1) and m.group(2):
        return m.group(1), m.group(2).strip()
    return "", str(s).strip()


rows = []
for _, r in td.iterrows():
    _c, full = parse_name(r["专业代号及名称"])
    core, note = split_core_note(full)
    sc, school = school_parts(r["院校代号及名称"])
    if school != "重庆外语外事学院":
        continue
    if not core.endswith("类"):
        continue
    rows.append({
        "sc": sc, "school": school, "full": full, "note": note,
        "xk": str(r["选考科目要求"]).strip(),
        "score": int(float(r["最低分"])), "rank": int(float(r["最低位次"])),
        "typ": str(r["专项计划/中外合作"]).strip() if pd.notna(r["专项计划/中外合作"]) else "普通",
    })

wb = Workbook()
ws = wb.active
ws.title = "大类查证清单"
headers = ["序号", "院校代码", "院校名称", "志愿类全名", "备注", "选科要求",
           "最低分", "最低位次", "招生类型", "查证状态", "包含专业(顿号分隔,待填)",
           "来源链接/年份(待填)", "备注(待填,如图片形式)"]
ws.append(headers)
hf = PatternFill("solid", fgColor="1D4ED8")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hf
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
for i, r in enumerate(rows, 1):
    ws.append([i, r["sc"], r["school"], r["full"], r["note"], r["xk"],
               r["score"], r["rank"], r["typ"], "待填", "", "", ""])
for c, w in enumerate([6, 9, 20, 34, 22, 11, 8, 9, 12, 9, 40, 26, 20], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
wb.save(OUT)
print("已生成:", OUT, " 行数:", len(rows))
for r in rows:
    print(f"  {r['sc']} {r['full']} ｜ {r['xk']} ｜ {r['score']}分/{r['rank']}名")
