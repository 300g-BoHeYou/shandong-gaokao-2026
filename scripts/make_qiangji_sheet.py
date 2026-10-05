# -*- coding: utf-8 -*-
"""
强基计划查证清单 + 网页数据生成（按用户 2026-04 反馈修订）

规则:
  1) 校考模式只说“出分前/出分后”，不使用民间俗称
  2) 仅面试院校 考查科目 标注“无笔试”
  3) 不展示院校级分数/位次，只保留分专业参考（分/位次）
  4) 强基专业名若为本科专业目录中的“具体专业”，直接预填到包含专业；
     大类名(如 化学类)保留待查；带括号方向的按去括号后名称预填并给出提示
  5) 来源冲突时以官方为准（东南=官网新闻，专业表为图片→留空；吉林=官方简章）
  6) 已确定专业标注软科 2026 排名状态
  7) 清华按文/理试验班标注分专业参考
输出:
  强基查证清单.xlsx (Sheet1 院校信息 / Sheet2 专业信息)
  data_qiangji.js
"""
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "2026山东高考第一批次投档.xlsx"
OUT_SHEET = ROOT / "强基查证清单.xlsx"
OUT_JS = ROOT / "data_qiangji.js"

# ============================================================
# 目录 & 软科
# ============================================================
with open(ROOT / "scripts" / "catalog_tree.json", encoding="utf-8") as f:
    cat = json.load(f)
cat_names, lei_names = set(), set()
for cls in cat["tree"].values():
    for lei, majors in cls.items():
        lei_names.add(lei)
        for _code, nm in majors:
            cat_names.add(nm)

with open(ROOT / "软科中国大学专业排名_2026.json", encoding="utf-8") as f:
    ruanke = json.load(f)
rk_major, name2keys, special_keys = {}, {}, set()
for men, cls in ruanke.items():
    for lei, majors in cls.items():
        for code_name, ranks in majors.items():
            name = code_name.split(" ", 1)[1] if " " in code_name else code_name
            rk_major[code_name] = dict(ranks)
            name2keys.setdefault(name, []).append(code_name)
            if ranks and set(ranks.values()) == {1}:
                special_keys.add(code_name)
rk_schools = {s for m in rk_major.values() for s in m}
rk_norm = {}
for s in rk_schools:
    rk_norm.setdefault(re.sub(r"\s+", "", s).replace("（", "(").replace("）", ")"), s)
SUFFIX = [r"克拉玛依校区$", r"秦皇岛分校$", r"威海分校$", r"青岛校区$", r"济南校区$",
          r"医学部$", r"医学院$", r"[（(][^（()）]*校区[）)]$", r"[（(][^（()）]*分校[）)]$",
          r"[（(][^（()）]*[）)]$"]


def canonical_school(name):
    n = re.sub(r"\s+", "", name).replace("（", "(").replace("）", ")")
    if n in rk_norm:
        return rk_norm[n]
    cur = n
    for _ in range(4):
        nxt = cur
        for pat in SUFFIX:
            n2 = re.sub(pat, "", nxt)
            if n2 != nxt and n2 in rk_norm:
                return rk_norm[n2]
        if nxt == cur:
            break
        cur = nxt
    return ""


def rank_info(school, major_name):
    """返回 {rk, st, cnt} — st: ranked/unranked/special/notpub/nosc"""
    rk_school = canonical_school(school)
    keys = name2keys.get(major_name) or []
    if not keys:
        return {"rk": None, "st": "notpub", "cnt": 0}
    key = keys[0]
    if key in special_keys:
        in_list = bool(rk_school and rk_school in rk_major[key])
        return {"rk": None, "st": "special", "cnt": len(rk_major[key]), "inList": in_list}
    ranks = rk_major[key]
    if not rk_school:
        return {"rk": None, "st": "nosc", "cnt": len(ranks)}   # 该校未收录软科(≠专业未排名)
    if rk_school in ranks:
        return {"rk": ranks[rk_school], "st": "ranked", "cnt": len(ranks)}
    return {"rk": None, "st": "unranked", "cnt": len(ranks)}


# ============================================================
# 投档表(分专业普通批参考)
# ============================================================
td = pd.read_excel(XLSX, sheet_name="第一次投档表", dtype=str)


def parse_name(name):
    m = re.match(r"^([0-9A-Za-z]{1,4})\s+(.*)$", str(name).strip())
    return (m.group(1), m.group(2).strip()) if m else ("", str(name).strip())


def split_core_note(full):
    m = re.search(r"[（(]", full)
    if m:
        return full[: m.start()].strip(), full[m.start() + 1:].rstrip(")）").strip()
    return full, ""


def school_parts(s):
    m = re.match(r"^([0-9A-Za-z]+)(.*)$", str(s).strip())
    if m and m.group(1) and m.group(2):
        return m.group(1), m.group(2).strip()
    return "", str(s).strip()


# 人工查证的试验班包含专业(可作为“包含该专业”的依据)
manual = {"试验班": {}}
mc_json = ROOT / "scripts" / "manual_check.json"
if mc_json.exists():
    with open(mc_json, encoding="utf-8") as f:
        manual = json.load(f)
syb_map = manual.get("试验班", {})

# 排除: 专项计划(已在类型列排除)、中外合作、分校/校区/马来西亚等异地校区、定向/预科/走读等特殊计划
EXCLUDE_PAT = re.compile(r"分校|校区|马来西亚|中外合作|合作办学|定向|预科|走读|联合培养|专项")

school_rows = {}   # 校 -> [(核心名, 分, 位次, 类型, 试验班查证专业, 全名)]
for _, r in td.iterrows():
    _c, full = parse_name(r["专业代号及名称"])
    core, note = split_core_note(full)
    sc, school = school_parts(r["院校代号及名称"])
    typ = str(r["专项计划/中外合作"]).strip() if pd.notna(r["专项计划/中外合作"]) else "普通"
    try:
        score, rank = int(float(r["最低分"])), int(float(r["最低位次"]))
    except (TypeError, ValueError):
        continue
    key = f"{sc}|{full}|{rank}"
    entry = syb_map.get(key) or {}
    syb_majors = entry.get("majors", []) if entry.get("mode") == "list" else []
    school_rows.setdefault(school, []).append((core, score, rank, typ, syb_majors, full))

# 大类名 -> 目录包含专业
lei2majors = {}
for cls in cat["tree"].values():
    for lei, majors in cls.items():
        lei2majors[lei] = [nm for _c, nm in majors]

# 强基专业类名 -> 可能的“具体专业名”列表(用于匹配同名志愿类或包含该专业的大类)
MAJOR_ALIAS = {
    "数学类": ["数学与应用数学", "信息与计算科学"],
    "物理学类": ["物理学", "应用物理学"],
    "化学类": ["化学", "应用化学"],
    "生物科学类": ["生物科学", "生物技术"],
    "历史学类": ["历史学"],
    "哲学类": ["哲学"],
    "力学类": ["理论与应用力学", "工程力学"],
    "核工程类": ["核工程与核技术", "工程物理"],
    "航空航天类": ["航空航天工程"],
    "中国语言文学类（古文字学方向）": ["汉语言文学"],
    "中国语言文学类（古文字学）": ["汉语言文学"],
    "汉语言文学（古文字学方向）": ["汉语言文学"],
    "汉语言文学（古文字学）": ["汉语言文学"],
    "汉语言文学（古文字方向）": ["汉语言文学"],
    "中国少数民族语言文学（古文字学方向）": ["中国少数民族语言文学"],
    "基础医学（八年制）": ["基础医学"],
    "草业科学（草类植物生物育种）": ["草业科学"],
    "药学（生物医药）": ["药学"],
    "测控技术与仪器（量子信息）": ["测控技术与仪器"],
    "复合材料与工程（航天材料）": ["复合材料与工程"],
    "材料科学与工程（航天材料）": ["材料科学与工程"],
}


def majors_of(cat):
    return MAJOR_ALIAS.get(cat, [cat])


def ref_best(school, cat):
    """分专业普通批参考: 同名志愿类 / 目录包含该专业的大类 / 已查证含该专业的试验班。
    仅取招生类型为“普通”(排除专项计划、中外合作办学)且非分校校区的志愿类,
    并在候选中取最低分(避免创新班等特殊班型分数偏高)。
    返回 (分, 位次, 采用的志愿类名)。"""
    cands = []
    for core, score, rank, typ, syb_majors, full in school_rows.get(school, []):
        if typ != "普通":
            continue
        if EXCLUDE_PAT.search(full):
            continue
        hit = False
        for mj in majors_of(cat):
            if core == mj or (core in lei2majors and mj in lei2majors[core]) or mj in syb_majors:
                hit = True
                break
        if hit:
            cands.append((score, rank, core))
    if not cands:
        return (None, None, "")
    score, rank, core = min(cands, key=lambda x: x[0])
    return (score, rank, core)


# 清华: 按文/理试验班标注
TSINGHUA_LI = ["数学与应用数学", "物理学", "化学", "生物科学", "信息与计算科学",
               "数理基础科学", "化学生物学", "理论与应用力学"]
TSINGHUA_WEN = ["中国语言文学类（古文字学方向）", "历史学类", "哲学类"]


def tsinghua_ref(cat):
    for core, score, rank, typ, syb_majors, full in school_rows.get("清华大学", []):
        if typ != "普通" or EXCLUDE_PAT.search(full):
            continue
        if cat in TSINGHUA_LI and core.startswith("理科试验班"):
            return (score, rank, core)
        if cat in TSINGHUA_WEN and core.startswith("文科试验班"):
            return (score, rank, core)
    return (None, None, "")


def derive_included(cat):
    """规则: 具体专业→直接预填; 带括号→去括号后判断, 括号内容作为备注; 大类→待查"""
    m = re.match(r"^(.*?)[（(](.*)[)）]$", cat)
    base = m.group(1).strip() if m else cat
    note = m.group(2).strip() if m else ""
    if base in cat_names and base not in lei_names:
        return [base], note
    return [], note


# ============================================================
# 院校数据
# ============================================================
SCHOOLS = [
    # ---------- 出分前校考 11 所 ----------
    dict(name="复旦大学", mode="出分前", examTime="初试6月12日，复试6月22日前", form="笔试+面试+体测",
         note="初试笔试科目按专业确定；A类考生校测成绩=初试100分+复试面试50分，B类考生为面试折算150分",
         majors=[("数学与应用数学", "数学"), ("信息与计算科学", "数学、计算科学"), ("物理学", "物理"),
                 ("化学", "化学"), ("生物科学类", "数理化综合"), ("基础医学", "数理化综合"),
                 ("中国语言文学类（古文字学）", "文史哲综合"), ("历史学", "文史哲综合"), ("哲学", "文史哲综合")]),
    dict(name="上海交通大学", mode="出分前", examTime="初试6月12日", form="笔试+面试+体测",
         note="校测分笔试(初试)、综合面试(复试)、体测；笔试科目按专业组区分",
         majors=[("数学与应用数学", "数学+物理"), ("物理学", "数学+物理"), ("生物医学工程", "数学+物理"),
                 ("工程力学", "数学+物理"), ("船舶与海洋工程", "数学+物理"), ("材料科学与工程", "数学+物理"),
                 ("化学", "数学+化学"), ("生物科学", "数学+化学"), ("生物医学科学", "数学+化学")]),
    dict(name="南京大学", mode="出分前", examTime="初试6月19日左右，复试6月21日左右", form="笔试+面试+体测",
         note="2026年变化：理科类初试考“数理探究”，文科类考“阅读表达”",
         majors=[("数学与应用数学", "数理探究"), ("信息与计算科学", "数理探究"), ("物理学", "数理探究"),
                 ("化学", "数理探究"), ("生物科学", "数理探究"), ("汉语言文学（古文字学）", "阅读表达"),
                 ("历史学", "阅读表达"), ("哲学", "阅读表达")]),
    dict(name="东南大学", mode="出分前", examTime="初试6月14日左右，复试6月18日左右",
         ratio="4倍（复试资格线）", form="笔试+面试+体测",
         entry="第一类考生参加初试(笔试)，按分省各专业计划4倍划定复试资格线；第二类考生(五科奥赛全国决赛二等奖及以上)不参加初试、直接入围复试",
         entryScore="初试成绩(笔试)", note="综合成绩=高考成绩÷高考满分×850+校测成绩；校测满分150分(初试100+复试50)；复试为体测+综合面试，无笔试",
         majors=[]),
    dict(name="中国科学技术大学", mode="出分前", examTime="初试6月14日上午8:30-11:50", form="笔试+面试+体测",
         note="一类生校考满分270分(笔试200+综合面试70)；二类生面试折算270分",
         majors=[("数学类", "数学、物理"), ("物理学类", "数学、物理"), ("化学", "数学、物理"),
                 ("生物科学类", "数学、物理"), ("理论与应用力学", "数学、物理"), ("核工程类", "数学、物理"),
                 ("地球物理学", "数学、物理"), ("量子信息科学", "数学、物理"), ("能源与动力工程", "数学、物理")]),
    dict(name="浙江大学", mode="出分前", examTime="初试6月16日18:30-21:00(紫金港)，体测6月19日，面试6月20日",
         form="笔试+面试+体测", note="2026年采用晚间笔试",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("工程力学", ""), ("化学", ""), ("生物工程", ""),
                 ("生物科学", ""), ("生态学", ""), ("生物育种科学", ""), ("基础医学", ""),
                 ("汉语言文学（古文字学方向）", ""), ("历史学", ""), ("哲学", "")]),
    dict(name="西安交通大学", mode="出分前", examTime="初试6月12日左右，复试6月18日左右", form="笔试+面试+体测",
         note="复试含笔试+面试，两环节各占综合成绩50%",
         majors=[("数学类", "数学、物理"), ("信息与计算科学", "数学、物理"), ("物理学", "数学、物理"),
                 ("生物技术", "数学、物理"), ("核工程与核技术", "数学、物理"), ("工程力学", "数学、物理"),
                 ("材料科学与工程", "数学、物理"), ("能源与动力工程", "数学、物理"),
                 ("储能科学与工程", "数学、物理"), ("哲学", "语文、历史")]),
    dict(name="同济大学", mode="出分前", examTime="初试6月13日14:00-16:30", form="笔试+面试+体测",
         note="学校考核满分150分：一类考生初试笔试50分+复试面试100分；初试为机考",
         majors=[("数学与应用数学", "数理综合"), ("应用物理学", "数理综合"), ("工程力学", "数理综合"),
                 ("海洋科学与技术", "数理综合"), ("应用化学", "化学"), ("生物技术", "化学")]),
    dict(name="厦门大学", mode="出分前", examTime="初试6月12日，复试6月17-18日", form="笔试+面试+体测",
         note="笔试100分、面试150分；2026年复试取消笔试(仅面试+体测)",
         majors=[("数学类", "数理综合"), ("物理学", "数理综合"), ("化学类", "数理化综合"),
                 ("生物科学类", "数理化综合"), ("海洋科学", "数理化综合"), ("药学", "数理化综合"),
                 ("历史学", "文史哲综合"), ("哲学", "文史哲综合")]),
    dict(name="北京航空航天大学", mode="出分前", examTime="初试6月12日，复试6月17-18日左右", form="笔试+面试+体测",
         note="笔试100分=数学及航空航天知识基础60分(数学基础50+航空航天10)+拔高科目40分；面试100分",
         majors=[("数学与应用数学", "数学拔高"), ("化学", "化学拔高"), ("信息与计算科学", "物理拔高"),
                 ("应用物理学", "物理拔高"), ("工程力学", "物理拔高"), ("测控技术与仪器（量子信息）", "物理拔高"),
                 ("材料科学与工程", "物理拔高"), ("航空航天类", "物理拔高")]),
    dict(name="中国人民大学", mode="出分前", examTime="初试6月13日左右，复试6月21日左右", form="笔试+面试+体测",
         note="一类考生初试60%+复试面试40%；A组考数理科技相关，B组考文史哲相关",
         majors=[("数据计算及应用", "数理科技相关"), ("汉语言文学（古文字学方向）", "文史哲相关"),
                 ("历史学类", "文史哲相关"), ("哲学", "文史哲相关")]),
    # ---------- 出分后校考 28 所 ----------
    dict(name="清华大学", mode="出分后", examTime="6月28日-7月3日综合考核", ratio="6倍",
         entry="高考成绩入围(6倍)；竞赛破格：五大学科奥赛全国决赛二等奖及以上",
         entryScore="高考成绩(不含政策加分)", form="笔试+面试+体测",
         signup="4月10日-4月30日(破格申请4月10-23日)",
         note="综合成绩满分100=高考折算85+笔试10+面试5；笔试为全国机考；分专业参考分按理科试验班/文科试验班标注",
         majors=[("数学与应用数学", "数学、物理、化学"), ("物理学", "数学、物理、化学"), ("化学", "数学、物理、化学"),
                 ("生物科学", "数学、物理、化学"), ("信息与计算科学", "数学、物理、化学"),
                 ("数理基础科学", "数学、物理、化学"), ("化学生物学", "数学、物理、化学"),
                 ("理论与应用力学", "数学、物理、化学"),
                 ("中国语言文学类（古文字学方向）", "数学、语文、历史"), ("历史学类", "数学、语文、历史"),
                 ("哲学类", "数学、语文、历史")]),
    dict(name="北京大学", mode="出分后", examTime="6月28日-7月4日考核", ratio="6倍",
         entry="高考成绩入围(6倍)；竞赛破格：五大学科奥赛全国决赛二等奖及以上",
         entryScore="高考成绩(不含政策加分)", form="笔试+面试+体测",
         signup="4月15日-4月30日(破格申请4月15-23日)",
         note="Ⅰ组/医学组笔试=数学、物理、化学；Ⅱ组笔试=语文、英语、历史",
         majors=[("数学类", "数学、物理、化学"), ("物理学类", "数学、物理、化学"), ("化学类", "数学、物理、化学"),
                 ("力学类", "数学、物理、化学"), ("生物科学类", "数学、物理、化学"),
                 ("基础医学（八年制）", "数学、物理、化学"),
                 ("历史学类", "语文、英语、历史"), ("考古学", "语文、英语、历史"),
                 ("哲学类", "语文、英语、历史"), ("中国语言文学类（古文字学方向）", "语文、英语、历史")]),
    dict(name="北京理工大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年取消竞赛破格；高分子材料与工程新增特种能源材料方向",
         majors=[("数学与应用数学", ""), ("应用物理学", ""), ("化学", ""), ("工程力学", ""),
                 ("智能无人系统技术", ""), ("高分子材料与工程", "")]),
    dict(name="中国农业大学", mode="出分后", form="", majors=[("生物科学", ""), ("生物育种科学", "")]),
    dict(name="北京师范大学", mode="出分后", form="", note="2026年取消竞赛破格入围通道",
         majors=[("汉语言文学（古文字学方向）", ""), ("历史学", ""), ("哲学", ""), ("数学与应用数学", ""),
                 ("物理学", ""), ("化学", ""), ("生物科学", "")]),
    dict(name="中央民族大学", mode="出分后", form="",
         majors=[("历史学", ""), ("哲学", ""), ("中国少数民族语言文学（古文字学方向）", "")]),
    dict(name="南开大学", mode="出分后", form="仅面试+体测",
         entry="高考成绩加权入围(理工类数学与物理或数学与化学1.5倍权重)+破格审核",
         note="加权示例：数学、物理按1.5倍计入入围成绩",
         majors=[("数学与应用数学", ""), ("信息与计算科学", ""), ("物理学", ""), ("化学", ""),
                 ("生物科学", ""), ("历史学", ""), ("哲学", ""), ("汉语言文学（古文字学方向）", "")]),
    dict(name="天津大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="船舶与海洋工程要求高考数学≥145分且加权成绩排序6倍以内；新增船舶与海洋工程；入围需达高考满分75%",
         majors=[("数学与应用数学", ""), ("应用物理学", ""), ("应用化学", ""), ("生物科学", ""),
                 ("工程力学", ""), ("合成生物学", ""), ("能源与动力工程", ""), ("船舶与海洋工程", "")]),
    dict(name="大连理工大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年无竞赛破格通道",
         majors=[("数学与应用数学", ""), ("信息与计算科学", ""), ("应用物理学", ""), ("工程力学", ""),
                 ("应用化学", ""), ("生物工程", "")]),
    dict(name="东北大学", mode="出分后", form="", majors=[("自动化", "")]),
    dict(name="吉林大学", mode="出分后", examTime="2026年6月30日开始（中心校区）", ratio="6倍",
         entry="高考成绩高于特招线80分(数理化)/60分(古文字学)基础上按6倍入围；数学与应用数学高考数学≥145分可破格入围",
         entryScore="高考成绩(不含政策加分)", form="笔试+面试+体测",
         signup="4月14日-4月30日(确认校考6月12-20日)",
         note="笔试、面试满分各100分，高校考核成绩=(笔试+面试)÷2；综合成绩=高考成绩÷7.5×85%+高校考核成绩×15%；数学与应用数学含“数学与应用数学、智能科学”两套培养方案",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("化学", ""), ("古文字学", "")]),
    dict(name="哈尔滨工业大学", mode="出分后", form="仅面试+体测",
         entry="高考成绩加权入围；第二类考生仅招金牌",
         note="第二类考生仅招收学科奥赛金牌；破格需审核",
         majors=[("智能装备与系统", ""), ("工程力学", ""), ("复合材料与工程（航天材料）", ""),
                 ("材料科学与工程（航天材料）", ""), ("飞行器制造工程", ""), ("储能科学与工程", ""),
                 ("核工程与核技术", ""), ("数学类", ""), ("应用物理学", "")]),
    dict(name="华东师范大学", mode="出分后", form="",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("生物科学", ""), ("哲学", ""),
                 ("汉语言文学（古文字学）", "")]),
    dict(name="山东大学", mode="出分后", form="", note="2026年新增密码科学与技术",
         majors=[("数学与应用数学", ""), ("密码科学与技术", ""), ("物理学", ""), ("化学", ""),
                 ("生物科学", ""), ("生物医学科学", ""), ("汉语言文学（古文字学方向）", ""),
                 ("历史学", ""), ("哲学", "")]),
    dict(name="中国海洋大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围",
         note="2026年无竞赛破格通道",
         majors=[("海洋科学", ""), ("海洋技术", ""), ("生物科学", "")]),
    dict(name="武汉大学", mode="出分后", form="",
         majors=[("哲学", ""), ("汉语言文学（古文字学方向）", ""), ("历史学", ""), ("数学与应用数学", ""),
                 ("物理学", ""), ("化学", ""), ("生物科学", ""), ("基础医学", ""), ("地球物理学", "")]),
    dict(name="华中科技大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年新增面试合格线",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("化学", ""), ("生物科学", ""), ("基础医学", ""),
                 ("汉语言文学（古文字学方向）", ""), ("哲学", "")]),
    dict(name="中南大学", mode="出分后", examTime="考核7月2日前完成", ratio="5倍",
         entry="高考成绩入围(5倍，专业志愿优先)；高考数学≥145分不受名额限制、满足第一志愿直接入围；竞赛破格：五大学科奥赛全国决赛二等奖及以上",
         entryScore="高考成绩(不含政策加分)", form="仅面试+体测",
         signup="4月10日-4月30日(确认考试6月12-20日)",
         note="面试满分100分、合格线60分；综合成绩=高考85%+面试15%；体育测试不合格不予录取",
         majors=[("数学与应用数学", ""), ("应用物理学", ""), ("应用化学", ""), ("生物科学", ""),
                 ("材料科学与工程", ""), ("信息与计算科学", "")]),
    dict(name="中山大学", mode="出分后", ratio="5倍", entry="高考成绩入围(5倍)+单科破格",
         note="6个专业可凭单科成绩破格：数学≥140(数学与应用数学/物理学/生态学/理论与应用力学)、物理≥100(物理学/理论与应用力学)、化学≥100(化学/药学（生物医药）)；2026年无笔试",
         form="仅面试+体测",
         majors=[("汉语言文学（古文字学方向）", ""), ("历史学", ""), ("哲学", ""), ("数学与应用数学", ""),
                 ("物理学", ""), ("化学", ""), ("生物科学", ""), ("生态学", ""), ("基础医学", ""),
                 ("理论与应用力学", ""), ("药学（生物医药）", "")]),
    dict(name="华南理工大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年取消竞赛破格；数学类入围成绩=高考成绩+数学×0.4、化学类/生物技术=高考成绩+化学×0.6、应用物理学=高考成绩+物理×0.6；数学类数学≥145可破格；面试低于60分不予录取",
         majors=[("数学类", ""), ("化学类", ""), ("生物技术", ""), ("应用物理学", "")]),
    dict(name="四川大学", mode="出分后", form="",
         majors=[("汉语言文学（古文字方向）", ""), ("历史学类", ""), ("哲学", ""), ("数学与应用数学", ""),
                 ("信息与计算科学", ""), ("物理学", ""), ("工程力学", ""), ("化学", ""),
                 ("生物科学", ""), ("基础医学", "")]),
    dict(name="重庆大学", mode="出分后", form="仅面试+体测", note="2026年无竞赛破格通道",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("储能科学与工程", "")]),
    dict(name="电子科技大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年无竞赛破格通道",
         majors=[("信息与计算科学", ""), ("数理基础科学", ""), ("应用物理学", "")]),
    dict(name="西北工业大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         majors=[("航空航天类", ""), ("船舶与海洋工程", ""), ("数学类", ""), ("应用物理学", ""), ("化学类", "")]),
    dict(name="西北农林科技大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年无竞赛破格通道",
         majors=[("生物育种科学", "")]),
    dict(name="兰州大学", mode="出分后", form="仅面试+体测", entry="高考成绩加权入围+单科破格",
         note="2026年无竞赛破格通道；新增核化工与核燃料工程",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("生物科学", ""), ("化学", ""),
                 ("草业科学（草类植物生物育种）", ""), ("生态学", ""), ("核化工与核燃料工程", ""),
                 ("历史学", ""), ("汉语言文学（古文字学方向）", "")]),
    dict(name="国防科技大学", mode="出分后", form="",
         majors=[("数学与应用数学", ""), ("物理学", ""), ("飞行器动力工程", "")]),
    dict(name="湖南大学", mode="出分后", form="仅面试+体测", note="2026年无竞赛破格通道",
         majors=[("化学", ""), ("应用化学", ""), ("化学生物学", "")]),
]

# ============================================================
# 官方招生章程链接与报名时间
# 来源: 江苏省教育考试院《2026年强基计划39校招生简章与报名时间汇总》
#       (扬州市教育局转载, http://jyj.yangzhou.gov.cn/jygl/zsks/ptgxzsxx/art/2026/art_49248a8b67e543beba06ca45a947f976.html)
# 链接指向各校官方微信公众号发布的招生简章; 未列出的备用官方入口为阳光高考强基专题
# ============================================================
CHSI_TOPIC = "https://gaokao.chsi.com.cn/gkzt/jcxkzs"
SCHOOL_LINKS = {
    "北京大学": "https://mp.weixin.qq.com/s?__biz=MzA3OTE0MjQzMw==&mid=2652034191&idx=1&sn=bb1331bbcd94c6fb515d71bf8515457e",
    "中国人民大学": "https://mp.weixin.qq.com/s?__biz=MzA3ODU4MzU3OA==&mid=2652504227&idx=1&sn=18c8b4d0139a0d6819ac09df0cab1e9a",
    "清华大学": "https://mp.weixin.qq.com/s?__biz=MzA4OTIyMzgxMw==&mid=2659401517&idx=1&sn=1a8939f9bb504acc6fe1f01e2a3079f8",
    "北京航空航天大学": "https://mp.weixin.qq.com/s?__biz=MzA3NTY1MTUzOQ==&mid=2651488755&idx=1&sn=78e3071a1058ac4f8bb53a4a307c3309",
    "北京理工大学": "https://mp.weixin.qq.com/s?__biz=MzUyMzYxNjE0Mg==&mid=2247653746&idx=1&sn=1b2c98c9a0fcdcd8ecdaea7c05639789",
    "中国农业大学": "https://mp.weixin.qq.com/s?__biz=MzA3MzM3MTUxOQ==&mid=2651285573&idx=1&sn=2e5c8920c491e9931c953eec00de74c6",
    "北京师范大学": "https://mp.weixin.qq.com/s?__biz=MjM5MTg4NjM2MQ==&mid=2658050606&idx=1&sn=95bf1324d037d15eddfcf05a359942d4",
    "中央民族大学": "https://mp.weixin.qq.com/s?__biz=MzAxNDEwNzkzMg==&mid=2651840949&idx=1&sn=86136bde471f66366ca8b07d79930ac3",
    "南开大学": "https://mp.weixin.qq.com/s?__biz=MjM5NzgwNDg0MQ==&mid=2650849867&idx=1&sn=4f30e34e45dd4b00753fb80124b7514b",
    "天津大学": "https://mp.weixin.qq.com/s?__biz=MjM5NzkwNzU0Mg==&mid=2654829695&idx=1&sn=0e93fc1c07ad21f95ae3306d0794714d",
    "大连理工大学": "https://mp.weixin.qq.com/s?__biz=MjM5MzIzODI2MA==&mid=2651909983&idx=1&sn=7200b71eb73ae90a0264ccb0b2378904",
    "东北大学": "https://mp.weixin.qq.com/s?__biz=MzA4NzIwNjIzMg==&mid=2651112701&idx=1&sn=b81ca43ffb1ff28440d95189f8627e68",
    "吉林大学": "https://mp.weixin.qq.com/s?__biz=MzUzODA1NjcyMA==&mid=2247577340&idx=1&sn=6c79ac646b148c11a0f12e7e5ef64d4f",
    "哈尔滨工业大学": "https://mp.weixin.qq.com/s?__biz=MzI3NzU5NTEwNw==&mid=2247700610&idx=1&sn=2b59d9cacfebcc9036743a817cb5a998",
    "复旦大学": "https://mp.weixin.qq.com/s?__biz=MjM5OTUwMzIwMQ==&mid=2651530990&idx=1&sn=3f3f32ee37f310c536ec9f5adfe4a272",
    "同济大学": "https://mp.weixin.qq.com/s?__biz=MzAxNzMxNzgxMQ==&mid=2656916233&idx=1&sn=b8e35f7b61328488758a36f45f810cb9",
    "上海交通大学": "https://mp.weixin.qq.com/s?__biz=MjM5MDIyMDQyMA==&mid=2650869498&idx=1&sn=3698662a6c0ed3146bfaab2994fe8d80",
    "华东师范大学": "https://mp.weixin.qq.com/s?__biz=MjM5MzY4NTMwMQ==&mid=2651604391&idx=1&sn=57159e05cb6c611183f75fb69fc9de8d",
    "南京大学": "https://mp.weixin.qq.com/s?__biz=MzAxODAzMjQ1NQ==&mid=2707754511&idx=1&sn=5ce19a08770ad4fd3da3d4400d2b0e5c",
    "东南大学": "https://mp.weixin.qq.com/s?__biz=MjM5NjQxMDE2MQ==&mid=2651046482&idx=1&sn=38a75c7039f270c856c30140ba209e2b",
    "浙江大学": "https://mp.weixin.qq.com/s?__biz=MjM5NDgxNTQwNQ==&mid=2651012201&idx=1&sn=6b3fda13c622514d801d73d475d7e9a9",
    "中国科学技术大学": "https://mp.weixin.qq.com/s?__biz=MjM5NTY5ODYyMQ==&mid=2652071906&idx=1&sn=ef3a752c953c911aac1442203ab5c726",
    "厦门大学": "https://mp.weixin.qq.com/s?__biz=MzA3OTM1MTIzNQ==&mid=2653260513&idx=1&sn=973f07612742b9bcbeebc727a775d1dc",
    "山东大学": "https://mp.weixin.qq.com/s?__biz=MzA4MTA0NTAyNw==&mid=2652300490&idx=1&sn=503462a8bc01ffa97680bf7c421f62ff",
    "中国海洋大学": "https://mp.weixin.qq.com/s?__biz=MjM5MjU2NDc0Mw==&mid=2651672109&idx=1&sn=6f96a006f49ea56db1565390c06f3b7a",
    "武汉大学": "https://mp.weixin.qq.com/s?__biz=MzA5MjM3MzUxMA==&mid=2656212784&idx=1&sn=ca2d83b88e002f39029c8d9c7bafb755",
    "华中科技大学": "https://mp.weixin.qq.com/s?__biz=MjM5MDAyMzM4MA==&mid=2650560908&idx=1&sn=2fd739c4893b4d80386bfad436425a38",
    "湖南大学": "https://mp.weixin.qq.com/s?__biz=MzA3ODQ3MjIzOA==&mid=2651823753&idx=1&sn=58545e3f871c8e38c0af60b835e1e24e",
    "中南大学": "https://mp.weixin.qq.com/s?__biz=MzA5NTI2MzExNA==&mid=2651242253&idx=1&sn=f938abb1099f1266580abe9b88140e30",
    "中山大学": "https://mp.weixin.qq.com/s?__biz=MzA3MDc5NjAyOA==&mid=2651597758&idx=1&sn=fa451b02ac3de5d82fe62ecc7d6b672d",
    "华南理工大学": "https://mp.weixin.qq.com/s?__biz=MzA5NDQ1NzE5Ng==&mid=2651568440&idx=1&sn=420a6750b04c4e41c8b5504ad542908b",
    "四川大学": "https://mp.weixin.qq.com/s?__biz=MzA5Mzc2MDYxMw==&mid=2650175673&idx=1&sn=405ab9a27b4a1833b5ac8f06f0455c57",
    "重庆大学": "https://mp.weixin.qq.com/s?__biz=MjM5MjA5MTQyMA==&mid=2651480156&idx=1&sn=02aed666cf288b4275e88c29fdcbf1d0",
    "电子科技大学": "https://mp.weixin.qq.com/s?__biz=MjM5MzM5MTMwMQ==&mid=2651769680&idx=1&sn=f672b72f879a82491df87a3fba93e305",
    "西安交通大学": "https://mp.weixin.qq.com/s?__biz=MjM5NzUxNjUyMQ==&mid=2650356875&idx=1&sn=89b41afe2cb0e47473c62c1e12689e4b",
    "西北工业大学": "https://mp.weixin.qq.com/s?__biz=MzA4NDE1MjgwNw==&mid=2649839945&idx=1&sn=49a6eee542f22e00d3d574865204139a",
    "西北农林科技大学": "https://mp.weixin.qq.com/s?__biz=MzA4MTE1MTkwNw==&mid=2651089853&idx=1&sn=7a3b96e7d715a4b5eee8c77a70213db4",
    "兰州大学": "https://mp.weixin.qq.com/s?__biz=MzA4OTQxMzUyNA==&mid=2651016404&idx=1&sn=8e7c8b262f6ff80d2869d11893e8ebc5",
    "国防科技大学": "https://mp.weixin.qq.com/s?__biz=MzAxNjYzNjkyNw==&mid=2654107743&idx=1&sn=2e8be4b6d2c80c1883bc8c12f3b07f23",
}
SCHOOL_SIGNUP = {
    "北京大学": "4月15日至30日；破格资格申请4月15日至23日",
    "中国人民大学": "第一类4月8日至30日；第二类4月8日至27日",
    "清华大学": "4月10日至30日；破格资格申请4月10日至23日",
    "北京航空航天大学": "4月13日至30日",
    "北京理工大学": "4月15日至30日",
    "中国农业大学": "4月7日至30日",
    "北京师范大学": "4月10日至30日",
    "中央民族大学": "4月10日至27日12:00",
    "南开大学": "第一类4月7日14:00至30日17:00；第二类4月7日14:00至29日17:00",
    "天津大学": "4月10日至30日",
    "大连理工大学": "4月7日至30日",
    "东北大学": "4月10日至30日",
    "吉林大学": "4月14日至30日",
    "哈尔滨工业大学": "第一类4月10日至30日；第二类4月10日至25日",
    "复旦大学": "A类4月8日至30日；B类4月8日至26日",
    "同济大学": "第一类4月9日至30日；第二类4月9日至26日",
    "上海交通大学": "第一类截至4月30日24:00；第二类截至4月27日24:00",
    "华东师范大学": "4月10日至30日15:00",
    "南京大学": "第一类4月10日至30日；第二类4月10日至26日",
    "东南大学": "4月8日至30日17:00",
    "浙江大学": "第一类4月13日至30日17:00；第二类4月13日至26日17:00",
    "中国科学技术大学": "第一类4月10日至30日；第二类4月10日至26日18:00",
    "厦门大学": "第一类4月10日至30日；第二类4月10日至26日",
    "山东大学": "4月10日至30日",
    "中国海洋大学": "4月9日至30日",
    "武汉大学": "4月10日12:00至30日12:00",
    "华中科技大学": "4月13日9:00至30日12:00",
    "湖南大学": "4月30日24:00前",
    "中南大学": "4月10日至30日",
    "中山大学": "4月10日至30日",
    "华南理工大学": "4月9日至30日",
    "四川大学": "第一类4月13日至30日；第二类4月13日至26日",
    "重庆大学": "4月10日至30日",
    "电子科技大学": "4月10日至30日",
    "西安交通大学": "4月10日至30日",
    "西北工业大学": "4月9日至30日",
    "西北农林科技大学": "4月8日至30日",
    "兰州大学": "4月10日至30日",
    "国防科技大学": "4月9日至30日",
}
for _s in SCHOOLS:
    _s["link"] = SCHOOL_LINKS.get(_s["name"], CHSI_TOPIC)
    if SCHOOL_SIGNUP.get(_s["name"]):
        _s["signup"] = SCHOOL_SIGNUP[_s["name"]]

# ============================================================
# 组装行
# ============================================================
rows = []
for s in SCHOOLS:
    no_exam = "仅面试" in (s.get("form") or "")
    if not s["majors"]:
        rows.append(dict(school=s["name"], cat="", inc="", incRk=[], subj="无笔试" if no_exam else "",
                         score=None, rank=None, ref="", note=""))
        continue
    for cat, subj in s["majors"]:
        inc_list, note = derive_included(cat)
        if s["name"] == "清华大学":
            sc, rk, ref = tsinghua_ref(cat)
        else:
            sc, rk, ref = ref_best(s["name"], cat)
        rows.append(dict(
            school=s["name"], cat=cat,
            inc="、".join(inc_list), incRk=[dict(n=m, **rank_info(s["name"], m)) for m in inc_list],
            subj=("无笔试" if no_exam else subj),
            score=sc, rank=rk, ref=ref, note=note,
        ))

# ============================================================
# Excel
# ============================================================
wb = Workbook()
ws1 = wb.active
ws1.title = "院校信息"
h1 = ["序号", "院校名称", "校考模式", "校考时间", "入围比例", "入围方式", "入围成绩计算",
      "校考形式", "竞赛破格政策", "报名时间", "备注"]
ws1.append(h1)
for i, s in enumerate(SCHOOLS, 1):
    ws1.append([i, s["name"], s.get("mode", ""), s.get("examTime", ""), s.get("ratio", ""),
                s.get("entry", ""), s.get("entryScore", ""), s.get("form", ""), s.get("brk", ""),
                s.get("signup", ""), s.get("note", "")])

ws2 = wb.create_sheet("专业信息")
h2 = ["序号", "院校名称", "强基专业类名称", "包含专业（顿号分隔）", "考查科目",
      "对应普通批最低分", "对应普通批最低位次", "参考志愿类", "备注（括号内方向）"]
ws2.append(h2)
for i, r in enumerate(rows, 1):
    ws2.append([i, r["school"], r["cat"], r["inc"], r["subj"],
                r["score"] if r["score"] else "", r["rank"] if r["rank"] else "", r["ref"], r["note"]])

for ws, headers, widths in [
    (ws1, h1, [5, 16, 9, 34, 14, 46, 26, 15, 34, 30, 60]),
    (ws2, h2, [5, 16, 30, 34, 18, 15, 17, 26, 22]),
]:
    fill = PatternFill("solid", fgColor="1D4ED8")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
wb.save(OUT_SHEET)
print("已生成:", OUT_SHEET)
print("  Sheet1 院校信息:", len(SCHOOLS), "行")
print("  Sheet2 专业信息:", len(rows), "行")

# ============================================================
# data_qiangji.js
# ============================================================
out_schools = []
ri = 0
for s in SCHOOLS:
    ms = []
    for cat, subj in s["majors"]:
        r = rows[ri]; ri += 1
        ms.append({"cat": r["cat"], "inc": r["inc"], "incRk": r["incRk"], "subj": r["subj"],
                   "s": r["score"], "r": r["rank"], "ref": r["ref"], "note": r["note"]})
    if not s["majors"]:
        r = rows[ri]; ri += 1
        ms.append({"cat": "", "inc": "", "incRk": [], "subj": r["subj"], "s": None, "r": None,
                   "ref": "", "note": ""})
    out_schools.append({
        "name": s["name"], "mode": s.get("mode", ""), "examTime": s.get("examTime", ""),
        "ratio": s.get("ratio", ""), "entry": s.get("entry", ""), "entryScore": s.get("entryScore", ""),
        "form": s.get("form", ""), "brk": s.get("brk", ""), "signup": s.get("signup", ""),
        "note": s.get("note", ""), "link": s.get("link", ""), "majors": ms,
    })

with open(OUT_JS, "w", encoding="utf-8") as f:
    f.write("/* 2026 强基计划数据(公开信息预填; “待查证”字段由人工查证后合并) */\n")
    f.write("window.QJ2026=" + json.dumps({"year": 2026, "schools": out_schools}, ensure_ascii=False) + ";\n")
print("已生成:", OUT_JS)
prefilled = sum(1 for r in rows if r["inc"])
matched = sum(1 for r in rows if r["score"])
print(f"  专业行 {len(rows)}: 预填包含专业 {prefilled} 条, 分专业参考分 {matched} 条, 考查科目 "
      f"{sum(1 for r in rows if r['subj'])} 条")
print("  待查专业类(空)院校:", [s["name"] for s in SCHOOLS if not s["majors"]])
