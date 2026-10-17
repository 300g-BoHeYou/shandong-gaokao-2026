# -*- coding: utf-8 -*-
"""生成 综招查证清单.xlsx: 19 所院校 × 关键字段 + 待查证事项清单。"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from zz_seed import PILOT          # noqa: E402
from zz_overseas import OVERSEAS   # noqa: E402

ALL = PILOT + OVERSEAS
HDR = ["序号", "院校名称", "范围", "院校层次", "2026计划(山东)", "报名时间",
       "校测日期", "校测环节", "校测形式", "校测地点", "入围比例", "综合成绩折算",
       "高考成绩要求", "招生专业(类)数", "官方简章链接", "备注/待核"]
FILL = PatternFill("solid", fgColor="EDE9FE")
HEAD = PatternFill("solid", fgColor="5B21B6")
WARN = PatternFill("solid", fgColor="FEF3C7")

wb = Workbook()
ws = wb.active
ws.title = "综招总表"
ws.append(HDR)
for c in range(1, len(HDR) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = HEAD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

todo = []
for i, s in enumerate(ALL, 1):
    evs = "、".join(f"{e['md']}·{e['label']}" + (f"({e['slot']})" if e.get("slot") else "")
                   for e in (s.get("events") or [])) or "待查证"
    nmaj = sum(len(g.get("majors") or []) for g in (s.get("groups") or []))
    gaps = []
    if s.get("plan") is None:
        gaps.append("山东计划数")
    if not (s.get("events") or []):
        gaps.append("校测日期")
    if any(m.get("plan") is None for g in (s.get("groups") or []) for m in (g.get("majors") or [])):
        gaps.append("分专业计划")
    row = [i, s["name"], "省外" if s.get("overseas") else "省内",
           s.get("level", ""), s.get("plan") if s.get("plan") is not None else "待查证",
           s.get("signup", ""), evs, s.get("examForm", ""), s.get("examDetail", "")[:60],
           s.get("examPlace", ""), s.get("ratio", ""), s.get("score", ""),
           s.get("gaokao", ""), nmaj, s.get("url", ""), "；".join(gaps) or ""]
    ws.append(row)
    if gaps:
        todo.append((s["name"], "；".join(gaps)))

widths = [5, 20, 7, 22, 12, 30, 30, 26, 40, 22, 20, 26, 30, 10, 46, 24]
for idx, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(idx)].width = w
for r in range(2, ws.max_row + 1):
    for c in range(1, len(HDR) + 1):
        ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 46
    if ws.cell(row=r, column=16).value:
        ws.cell(row=r, column=16).fill = WARN
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HDR))}{ws.max_row}"

# 第二表: 待查证事项
ws2 = wb.create_sheet("待查证事项")
ws2.append(["序号", "院校", "待补内容", "建议来源"])
for c in range(1, 5):
    cell = ws2.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = HEAD
for i, (name, g) in enumerate(todo, 1):
    src = "学校本科招生网 / 山东省教育招生考试院《填报志愿指南》"
    ws2.append([i, name, g, src])
ws2.column_dimensions["A"].width = 5
ws2.column_dimensions["B"].width = 22
ws2.column_dimensions["C"].width = 28
ws2.column_dimensions["D"].width = 46
for r in range(2, ws2.max_row + 1):
    ws2.row_dimensions[r].height = 28

out = ROOT / "综招查证清单.xlsx"
wb.save(out)
print(f"已生成 {out.name}: 19 所（省内10/省外9）")
print(f"待查证事项 {len(todo)} 条:")
for n, g in todo:
    print(f"  {n}: {g}")
