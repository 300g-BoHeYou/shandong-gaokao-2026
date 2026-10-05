# -*- coding: utf-8 -*-
"""把 39 校官方简章抽取结果合并进 强基查证清单.xlsx (官方数据覆盖人工填写)。"""
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "qj_cache"
XLSX = ROOT / "强基查证清单.xlsx"

FINAL = json.loads((CACHE / "final_fields.json").read_text(encoding="utf-8"))
CORPUS = json.loads((CACHE / "corpus.json").read_text(encoding="utf-8"))
CN = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
ALT = {"中央民族大学": "fix_69437_600.html", "上海交通大学": "fix_38987_52.html",
       "四川大学": "fix_scu_jz.html", "兰州大学": "lzu_jz.html"}


def to_text(html):
    h = re.sub(r"<script.*?</script>|<style.*?</style>|<!--.*?-->", " ", html, flags=re.S | re.I)
    h = re.sub(r"<br\s*/?>|</p>|</div>|</tr>|</li>|</h[1-6]>|</table>", "\n", h, flags=re.I)
    h = re.sub(r"</t[dh]>", "\t", h, flags=re.I)
    h = re.sub(r"<[^>]+>", "", h)
    h = (h.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">"))
    return re.sub(r"\s+", "", h)


def text_of(name):
    t = CORPUS.get(name, {}).get("text", "")
    if len(t) < 1200 and name in ALT:
        p = CACHE / ALT[name]
        if p.exists():
            t = to_text(p.read_text(encoding="utf-8", errors="replace"))
    return re.sub(r"\s+", "", t)


def ratio_of(name, t):
    """入围比例: 取含'入围'句中第一个非破格语境的 N倍"""
    for s in re.split(r"(?<=[。；])", t):
        if "入围" in s and "倍" in s:
            for m in re.finditer(r"([0-9一二三四五六七八九])\s*倍", s):
                if "破格" not in s[max(0, m.start() - 22):m.start()]:
                    v = m.group(1)
                    return (str(CN[v]) if v in CN else v) + "倍"
    for s in re.split(r"(?<=[。；])", t):
        if "倍" in s and ("合格线" in s or "计划数" in s):
            m = re.search(r"([0-9一二三四五六七八九])\s*倍", s)
            if m:
                v = m.group(1)
                return (str(CN[v]) if v in CN else v) + "倍"
    return ""


def form_of(t):
    no_exam = any(k in t for k in ["不设笔试", "无笔试", "取消笔试", "不再组织笔试", "不进行笔试", "仅有面试", "仅面试"])
    has_bs = ("笔试" in t or "初试" in t) and not no_exam
    has_ms = "面试" in t
    has_tc = ("体育测试" in t) or ("体测" in t)
    if has_bs:
        return "+".join([x for x, c in [("笔试", True), ("面试", has_ms), ("体测", has_tc)] if c])
    return "仅面试" + ("+体测" if has_tc else "")


def win_of(name, t):
    """校考时间: 分别抓 初试/复试/考核 的日期, 组合成简明串"""
    parts = []
    for key, label in [("初试", "初试"), ("复试", "复试")]:
        best = ""
        for s in re.split(r"(?<=[。；])", t):
            if key in s:
                m = re.search(r"([0-9]{1,2}月[0-9]{1,2}日)(前|左右|至[0-9]{1,2}月[0-9]{1,2}日)?", s)
                if m:
                    suf = ""
                    if "前" in s[max(0, m.start() - 6):m.end() + 2]:
                        suf = "前"
                    elif "左右" in s[max(0, m.start() - 6):m.end() + 4]:
                        suf = "左右"
                    best = f"{label}{m.group(1)}{suf}"
                    break
        if best:
            parts.append(best)
    if not parts:
        for s in re.split(r"(?<=[。；])", t):
            if any(k in s for k in ["考核时间", "测试时间", "校测时间", "考核安排"]):
                m = re.search(r"([0-9]{1,2}月[0-9]{1,2}日)\s*[至到\-—~－]\s*([0-9]{1,2}月[0-9]{1,2}日)", s)
                if m:
                    parts.append(f"{m.group(1)}-{m.group(2)}")
                    break
                m = re.search(r"([0-9]{1,2}月[0-9]{1,2}日)", s)
                if m:
                    parts.append(m.group(1))
                    break
    if not parts:
        f = FINAL.get(name, {})
        ds = [d for d in f.get("dates", []) if d.startswith(("6月", "7月"))]
        if ds:
            parts.append("、".join(ds[:3]))
    return "；".join(parts)


def brk_of(t):
    """竞赛破格政策: 只保留真正讲破格资格的句子"""
    KW = ("第二类", "竞赛", "奥赛", "奥林匹克", "奖项", "金牌", "银牌", "资格申请")
    sents = re.split(r"(?<=[。；])", t)
    for s in sents:
        if "破格" in s and any(k in s for k in KW) and any(k in s for k in ["取消", "不再设置", "不再有"]):
            return s[:110]
    for s in sents:
        if "破格" in s and any(k in s for k in KW) and any(k in s for k in ["二等奖", "一等奖", "金牌", "银牌", "全国决赛"]):
            return s[:110]
    for s in sents:
        if "破格" in s and any(k in s for k in KW) and "审核" in s:
            return s[:110]
    for s in sents:
        if "破格" in s and any(k in s for k in KW):
            return s[:110]
    return ""


def entry_of(t):
    main = ""
    for s in re.split(r"(?<=[。；])", t):
        if "入围" in s and any(k in s for k in ["按照", "依据", "确定", "不低于", "达到"]):
            main = s[:110]
            break
    extra = ""
    m = re.search(r"[^。；]{0,30}(?:145分|单科成绩特别优秀|单科成绩)[^。；]{0,70}", t)
    if m:
        extra = m.group(0)
    if extra and extra not in main:
        return (main + "；" + extra)[:190]
    return main


def calc_of(t):
    for s in re.split(r"(?<=[。；])", t):
        if any(k in s for k in ["加权", "单科成绩", "入围成绩"]) and "成绩" in s:
            return s[:110]
    return ""


def subj_of(t):
    subs = []
    for s in re.split(r"(?<=[。；])", t):
        if "笔试" in s and any(k in s for k in ["科目", "考查", "考察", "内容"]):
            for x in ["数学", "物理", "化学", "生物", "语文", "历史", "哲学", "英语", "古文字"]:
                if x in s and x not in subs:
                    subs.append(x)
    return "、".join(subs)


def main():
    out_path = ROOT / "强基查证清单_官方核对.xlsx"
    wb = load_workbook(XLSX)
    ws1, ws2 = wb["院校信息"], wb["专业信息"]
    hdr1 = [c.value for c in ws1[1]]
    ix = {h: i + 1 for i, h in enumerate(hdr1)}

    changes = []
    for r in range(2, ws1.max_row + 1):
        name = ws1.cell(row=r, column=ix["院校名称"]).value
        if not name or name not in FINAL:
            continue
        t = text_of(name)
        subs = subj_of(t)
        vals = {
            "入围比例": ratio_of(name, t),
            "入围方式": entry_of(t),
            "入围成绩计算": calc_of(t),
            "校考时间": win_of(name, t),
            "校考形式": form_of(t),
            "竞赛破格政策": brk_of(t),
            "报名时间": (CORPUS.get(name, {}).get("signup", "") or "").replace(" 至 ", "至"),
        }
        for k, v in vals.items():
            if not v:
                continue
            old = str(ws1.cell(row=r, column=ix[k]).value or "").strip()
            # 入围比例: 旧值含同样倍数时保留更详细的旧值
            if k == "入围比例" and v.rstrip("倍") in old:
                continue
            if old != v:
                ws1.cell(row=r, column=ix[k]).value = v
                changes.append((name, k, old[:26], v[:44]))

    print("Sheet1 更新:", len(changes), "处")
    for c in changes:
        print("   ", c[0], "|", c[1], "|", c[2], "->", c[3])
    wb.save(out_path)
    print("\n已保存(原文件被 Excel 占用):", out_path)


if __name__ == "__main__":
    main()
